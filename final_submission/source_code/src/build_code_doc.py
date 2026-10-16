import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from code_clean import strip_python, strip_requirements

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
TESTS = os.path.join(ROOT, "tests")
CONFIG = os.path.join(ROOT, "source_code", "config.json")
REQ = os.path.join(ROOT, "source_code", "requirements.txt")
OUT = os.path.join(ROOT, "paper")
NAME = "B106_Code"
BODY, MONO = "Calibri", "Consolas"
NAVY = RGBColor(0x14, 0x21, 0x3D)
GREY = RGBColor(0x55, 0x55, 0x55)

PARTS = [
    ("Core algorithm", "The watermarking method itself: embedding, extraction, metrics, attacks, logo and image loading.", [
        ("src", "watermark_core.py", "Hybrid DWT, DCT and SVD embedding and extraction"),
        ("src", "metrics.py", "PSNR, SSIM, MSE, NC, ZNCC and BER"),
        ("src", "attacks.py", "Fourteen attacks plus the unattacked reference"),
        ("src", "logo.py", "Procedural 64x64 binary watermark logo"),
        ("src", "image_io.py", "Safe image loading and standardisation"),
    ]),
    ("Dataset", "Everything needed to rebuild the 100 image evaluation set.", [
        ("src", "dataset_manifest.py", "List of source images, licences and categories"),
        ("src", "build_dataset.py", "Downloads, standardises and records the 100 images"),
    ]),
    ("Experiments", "Scripts that run the experiments and record results.", [
        ("src", "run_experiment.py", "Original two image experiment"),
        ("src", "run_batch_experiment.py", "Batch experiment on 100 images"),
        ("src", "video_demo.py", "Per frame video test"),
        ("src", "exp_jpeg_truncation.py", "Truncation versus rounding before JPEG"),
        ("src", "exp_principal_component_variant.py", "Principal component embedding variant"),
    ]),
    ("Analysis and benchmarks", "Statistics, figures and comparisons built from the recorded results.", [
        ("src", "analysis.py", "Tables and figures for the main results"),
        ("src", "analysis_tech.py", "Figures for the technology comparison"),
        ("src", "benchmark_tech.py", "Library and scheme benchmarks"),
        ("src", "make_diagram.py", "Pipeline diagrams"),
        ("src", "make_paper_figures.py", "Column width figures for the paper"),
    ]),
    ("Document generators", "Scripts that assemble the report, presentation and paper from the results.", [
        ("src", "docx_helpers.py", "Helper functions for Word documents"),
        ("src", "build_report.py", "Builds the project report"),
        ("src", "build_deck.py", "Builds the presentation"),
        ("src", "build_paper.py", "Builds the research paper"),
        ("src", "validate_submission.py", "Checks the finished package"),
    ]),
    ("Tests", "Automated tests (pytest).", [
        ("tests", "conftest.py", "Test configuration"),
        ("tests", "test_watermark_core.py", "Tests for embedding and extraction"),
        ("tests", "test_metrics.py", "Tests for the metrics"),
        ("tests", "test_attacks.py", "Tests for the attacks"),
        ("tests", "test_image_io.py", "Tests for image loading"),
        ("tests", "test_dataset_and_results.py", "Checks on the dataset and results"),
    ]),
    ("Configuration", "Settings and dependencies.", [
        ("cfg", "config.json", "Experiment settings"),
        ("cfg", "requirements.txt", "Python packages used"),
    ]),
]


def read_clean(kind, name):
    if kind == "src":
        path = os.path.join(HERE, name)
    elif kind == "tests":
        path = os.path.join(TESTS, name)
    else:
        path = CONFIG if name == "config.json" else REQ
    text = open(path, encoding="utf-8").read()
    if name.endswith(".py"):
        return strip_python(text)
    if name.endswith(".txt"):
        return strip_requirements(text)
    return text.rstrip("\n") + "\n"


def field(run, instr):
    for t, txt in (("begin", None), (None, instr), ("separate", None), (None, "1"), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
            run._r.append(el)
        elif txt == instr:
            it = OxmlElement("w:instrText")
            it.set(qn("xml:space"), "preserve")
            it.text = instr
            run._r.append(it)
        else:
            tt = OxmlElement("w:t")
            tt.text = txt
            run._r.append(tt)


def shade(par, fill):
    pPr = par._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def build(pages):
    doc = Document()
    sec = doc.sections[0]
    sec.page_height, sec.page_width = Cm(29.7), Cm(21.0)
    sec.left_margin = sec.right_margin = Cm(1.8)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    sec.different_first_page_header_footer = True
    sec.first_page_header.is_linked_to_previous = False
    sec.first_page_footer.is_linked_to_previous = False
    st = doc.styles
    st["Normal"].font.name = BODY
    st["Normal"].font.size = Pt(11)
    st["Normal"].element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    st["Normal"].paragraph_format.space_after = Pt(4)
    for name, size, color, before, after in (("Heading 1", 18, NAVY, 0, 6), ("Heading 2", 12.5, NAVY, 14, 3)):
        h = st[name]
        h.font.name = BODY
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn("w:ascii"), BODY)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), BODY)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run()
    r.font.size = Pt(9)
    r.font.color.rgb = GREY
    field(r, "PAGE")
    hp = sec.header.paragraphs[0]
    hr = hp.add_run("Digital Image Watermarking, project source code   |   Roll No. B106")
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = GREY

    def text(t, size=11, bold=False, italic=False, align=None, after=4, color=None):
        p = doc.add_paragraph()
        if align == "center":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        run = p.add_run(t)
        run.font.size = Pt(size)
        run.bold, run.italic = bold, italic
        if color is not None:
            run.font.color.rgb = color
        return p

    for _ in range(5):
        doc.add_paragraph()
    text("Digital Image Watermarking", 28, True, align="center", after=2, color=NAVY)
    text("Project source code", 16, align="center", after=14)
    text("A hybrid DWT, DCT and SVD approach", 12, italic=True, align="center", after=40)
    text("Krishang Vikram Singh", 13, True, align="center", after=1)
    text("Roll No. B106   Batch C1", 11, align="center", after=1)
    text("B.Tech Computer Engineering, NMIMS MPSTME", 11, align="center", after=1)
    text("Subject: Image and Video Processing", 11, align="center", after=40)
    text("Language: Python 3.13     Main libraries: NumPy, SciPy, scikit-image, PyWavelets, Pillow, pandas, matplotlib", 9.5, color=GREY, align="center", after=1)
    text("Optional: OpenCV and FFmpeg (video test), blind-watermark (comparison)", 9.5, color=GREY, align="center", after=0)
    doc.add_page_break()

    entries = []
    for pi, (ptitle, pdesc, files) in enumerate(PARTS, 1):
        entries.append((0, f"Part {pi}: {ptitle}", f"Part {pi}: {ptitle}"))
        for fi, (kind, fname, desc) in enumerate(files, 1):
            label = f"{pi}.{fi}  {fname}"
            entries.append((1, label, label))
    doc.add_heading("Contents", level=1)
    for lvl, label, key in entries:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.left_indent = Cm(0.8 * lvl)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(17.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        run = p.add_run(f"{label}\t{pages.get(key, '0')}")
        run.font.size = Pt(10.5 if lvl == 0 else 10)
        run.bold = lvl == 0

    doc.add_page_break()
    doc.add_heading("How the project is organised", level=1)
    text("The code is grouped into seven parts. The core algorithm is small and has no dependency on the other parts. "
         "The dataset builder prepares the images, the experiment scripts run the algorithm on them and record results, "
         "and the analysis scripts turn the recorded results into tables and figures. The document generators assemble the report, "
         "presentation and paper from those results. Tests check the algorithm, the loaders and the completeness of the results.", 11, after=8)
    tbl = doc.add_table(rows=1, cols=3)
    tbl.style = "Table Grid"
    for i, h in enumerate(("File", "Part", "Purpose")):
        c = tbl.rows[0].cells[i]
        c.text = ""
        run = c.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        tcPr = c._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:color"), "auto")
        shd.set(qn("w:fill"), "14213D")
        tcPr.append(shd)
    for pi, (ptitle, pdesc, files) in enumerate(PARTS, 1):
        for fi, (kind, fname, desc) in enumerate(files, 1):
            cells = tbl.add_row().cells
            for c, v in zip(cells, (fname, f"{pi}. {ptitle}", desc)):
                c.text = ""
                run = c.paragraphs[0].add_run(v)
                run.font.size = Pt(9)
                c.paragraphs[0].paragraph_format.space_after = Pt(0)
    for row in tbl.rows:
        for c, w in zip(row.cells, (5.6, 4.4, 7.2)):
            c.width = Cm(w)
    text("", after=6)
    text("Order in which the main scripts are run", 12, True, after=3, color=NAVY)
    for i, step in enumerate(("build_dataset.py: downloads and standardises the 100 images and writes the metadata table.",
                              "run_batch_experiment.py: embeds the watermark, applies the attacks, extracts and records every result.",
                              "analysis.py: statistics, tables and figures from the recorded results.",
                              "benchmark_tech.py and analysis_tech.py: library and scheme comparisons.",
                              "video_demo.py: optional per frame video test.",
                              "pytest tests: runs the automated tests."), 1):
        p = text(f"{i}.  {step}", 10.5, after=2)
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.6)

    total = 0
    for pi, (ptitle, pdesc, files) in enumerate(PARTS, 1):
        h = doc.add_heading(f"Part {pi}: {ptitle}", level=1)
        h.paragraph_format.page_break_before = True
        text(pdesc, 11, italic=True, after=6)
        for fi, (kind, fname, desc) in enumerate(files, 1):
            code = read_clean(kind, fname)
            lines = code.rstrip("\n").split("\n")
            total += len(lines)
            hh = doc.add_heading(f"{pi}.{fi}  {fname}", level=2)
            if fi > 1:
                hh.paragraph_format.page_break_before = False
            text(f"{desc}. {len(lines)} lines.", 10, italic=True, color=GREY, after=4)
            for line in lines:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(0.15)
                run = p.add_run(line.replace("\t", "    ") if line else " ")
                run.font.name = MONO
                run._r.rPr.rFonts.set(qn("w:eastAsia"), MONO)
                run.font.size = Pt(8)
                shade(p, "F5F5F3")
    doc.core_properties.title = "B106 source code"
    doc.core_properties.author = "Krishang Vikram Singh"
    path = os.path.join(OUT, NAME + ".docx")
    doc.save(path)
    return path, entries, total


def find_pages(pdf, entries):
    n = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True, text=True).stdout).group(1))
    texts = [subprocess.run(["pdftotext", "-layout", "-f", str(i), "-l", str(i), pdf, "-"], capture_output=True, text=True).stdout for i in range(1, n + 1)]
    start = next(i for i, t in enumerate(texts) if "How the project is organised" in t) + 1
    pages = {}
    for _, label, key in entries:
        for i in range(start, n):
            norm = re.sub(r"\s+", " ", label)
            if any(re.sub(r"\s+", " ", line.strip()).startswith(norm) and "...." not in line for line in texts[i].splitlines()):
                pages[key] = i + 1
                break
    return pages, n


if __name__ == "__main__":
    os.makedirs(OUT, exist_ok=True)
    path, entries, total = build({})
    pages = {}
    for _ in range(2):
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, path], capture_output=True, timeout=600)
        pages, n = find_pages(os.path.join(OUT, NAME + ".pdf"), entries)
        path, entries, total = build(pages)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", OUT, path], capture_output=True, timeout=600)
    print("saved", path, "| code lines:", total, "| pages:", n, "| unresolved:", [k for _, _, k in entries if k not in pages])
