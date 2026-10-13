"""
build_deck.py
=============
Builds presentation/Final_Presentation.pptx in the visual language of the first submission
(Calibri, navy/teal palette, numbered kicker + bold title, tile cards) from the recorded
results, and writes presentation/SPEAKER_NOTES.md.

    python build_deck.py
"""

from __future__ import annotations

import json
import os
import re

import pandas as pd
from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
RES, VIS, TAB = os.path.join(ROOT, "results"), os.path.join(ROOT, "visualizations"), os.path.join(ROOT, "results", "tables")
OUT = os.path.join(ROOT, "presentation")
os.makedirs(OUT, exist_ok=True)
V = lambda n: os.path.join(VIS, n)

AS = json.load(open(os.path.join(RES, "analysis_summary.json")))
RS = json.load(open(os.path.join(RES, "run_summary.json")))
TS = json.load(open(os.path.join(RES, "tech_summary.json")))
ROB = pd.read_csv(os.path.join(TAB, "robustness_by_attack.csv")).set_index("attack")
SWP = pd.read_csv(os.path.join(TAB, "alpha_sweep_summary.csv")).set_index(["attack", "alpha"])
CTL = pd.read_csv(os.path.join(TAB, "controls_summary.csv"))
COR = pd.read_csv(os.path.join(TAB, "nc_vs_image_statistics.csv")).set_index("attack")
TSCH = pd.read_csv(os.path.join(RES, "tech_benchmark_schemes.csv")).set_index(["scheme", "attack"])
VID = pd.read_csv(os.path.join(RES, "video_demo_results.csv"))
META = pd.read_csv(os.path.join(ROOT, "data", "metadata.csv"))
E1, E2, E4, BASE = AS["E1"], AS["E2"], AS["E4"], AS["metric_baselines"]
G = CTL.iloc[0]
CA, CB = CTL[CTL.control.str.startswith("a:")].iloc[0], CTL[CTL.control.str.startswith("b:")].iloc[0]
CC = CTL[CTL.control.str.startswith("c: ")].iloc[0]
PROJ, BWM, BDCT, LSB = ("Project: global DWT-DCT-SVD", "blind-watermark 0.4.4 (block DWT-DCT-SVD)", "Baseline: 8x8 block DCT", "Baseline: spatial LSB")
ber = lambda s, a: float(TSCH.loc[(s, a), "mean_ber"])
psnr_s = lambda s: float(TSCH.reset_index().query("scheme == @s").psnr_db.mean())
tms = lambda s, c: float(TSCH.reset_index().query("scheme == @s")[c].median())
sw = lambda a, al, c: float(SWP.loc[(a, al), c])
pc = lambda x, d=0: f"{100 * x:.{d}f} %"

# ------------------------------------------------------------------ design tokens (from the first deck)
NAVY, NAVY2, NAVY3 = RGBColor(0x13, 0x23, 0x47), RGBColor(0x0C, 0x18, 0x30), RGBColor(0x1B, 0x2E, 0x57)
TEAL, TEALD = RGBColor(0x17, 0xB8, 0xB0), RGBColor(0x0E, 0x8F, 0x89)
INK, MUTED, LIGHT, LINE = RGBColor(0x1A, 0x22, 0x33), RGBColor(0x66, 0x70, 0x85), RGBColor(0xF3, 0xF5, 0xF9), RGBColor(0xDC, 0xE2, 0xEB)
WHITE, PALE, RED, REDBG, TEALBG = RGBColor(0xFF, 0xFF, 0xFF), RGBColor(0x93, 0xA8, 0xC7), RGBColor(0xC1, 0x44, 0x3C), RGBColor(0xFB, 0xEC, 0xEA), RGBColor(0xE6, 0xF6, 0xF5)
FONT = "Calibri"

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
BLANK = prs.slide_layouts[6]
NOTES = []


def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=0.75, radius=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid(); sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line; sh.line.width = Pt(lw)
    sh.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = radius
    return sh


def text(s, x, y, w, h, content, size=12, color=INK, bold=False, align=None, anchor=MSO_ANCHOR.TOP, italic=False, spacing=None, line_sp=None):
    """content: str or list of str (paragraphs). Inline **bold** and `code`."""
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = anchor
    paras = content if isinstance(content, list) else [content]
    for i, ptxt in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        if align:
            p.alignment = align
        if spacing is not None:
            p.space_after = Pt(spacing)
        if line_sp:
            p.line_spacing = line_sp
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", ptxt):
            if not part:
                continue
            r = p.add_run()
            if part.startswith("**"):
                r.text = part[2:-2]; r.font.bold = True
            elif part.startswith("`"):
                r.text = part[1:-1]; r.font.name = "Consolas"; r.font.bold = bold
                r.font.size = Pt(size - 1); r.font.color.rgb = TEALD if color == INK else color
                continue
            else:
                r.text = part; r.font.bold = bold
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.italic = italic
    return tb


def bullets(s, x, y, w, h, items, size=12.5, color=INK, gap=5, dot=TEAL):
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.04)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(gap)
        r0 = p.add_run(); r0.text = "▪  "; r0.font.size = Pt(size); r0.font.color.rgb = dot; r0.font.name = FONT
        for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", it):
            if not part:
                continue
            r = p.add_run()
            if part.startswith("**"):
                r.text = part[2:-2]; r.font.bold = True
            elif part.startswith("`"):
                r.text = part[1:-1]; r.font.name = "Consolas"; r.font.size = Pt(size - 1); r.font.color.rgb = TEALD
                continue
            else:
                r.text = part
            r.font.name = FONT; r.font.size = Pt(size); r.font.color.rgb = color
    return tb


def picture(s, path, x, y, w=None, h=None):
    """Fit the picture into the w x h box (inches), centred, preserving aspect."""
    iw, ih = Image.open(path).size
    ar = iw / ih
    if w is not None and h is not None:
        if w / h > ar:
            pw, ph = h * ar, h
        else:
            pw, ph = w, w / ar
        px, py = x + (w - pw) / 2, y + (h - ph) / 2
    elif w is not None:
        pw, ph, px, py = w, w / ar, x, y
    else:
        pw, ph, px, py = h * ar, h, x, y
    return s.shapes.add_picture(path, Inches(px), Inches(py), Inches(pw), Inches(ph))


def tile(s, x, y, w, h, big, label, fill=NAVY, bigc=TEAL, labc=PALE, bigsize=28):
    rect(s, x, y, w, h, fill=fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, x + 0.15, y + 0.1, w - 0.3, h * 0.55, big, size=bigsize, color=bigc, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.15, y + h * 0.58, w - 0.3, h * 0.4, label, size=10.5, color=labc)


def card(s, x, y, w, h, title, body, fill=LIGHT, line=LINE, tcolor=NAVY, size=11):
    rect(s, x, y, w, h, fill=fill, line=line, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.04)
    text(s, x + 0.15, y + 0.1, w - 0.3, 0.35, title, size=12.5, color=tcolor, bold=True)
    if isinstance(body, list):
        bullets(s, x + 0.12, y + 0.5, w - 0.25, h - 0.55, body, size=size, gap=3)
    else:
        text(s, x + 0.15, y + 0.5, w - 0.3, h - 0.55, body, size=size, color=INK)


def table(s, x, y, w, rows, header, widths, size=10.5, row_h=0.34, header_fill=NAVY, zebra=True, bold_first=False):
    nrows, ncols = len(rows) + 1, len(header)
    shp = s.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
    tb = shp.table
    tot = sum(widths)
    for i, wd in enumerate(widths):
        tb.columns[i].width = Inches(w * wd / tot)
    for r_i in range(nrows):
        tb.rows[r_i].height = Inches(row_h)
        for c_i in range(ncols):
            cell = tb.cell(r_i, c_i)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.025)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            val = header[c_i] if r_i == 0 else rows[r_i - 1][c_i]
            tf = cell.text_frame; tf.word_wrap = True
            p = tf.paragraphs[0]
            for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", str(val)):
                if not part:
                    continue
                r = p.add_run()
                if part.startswith("**"):
                    r.text = part[2:-2]; r.font.bold = True
                elif part.startswith("`"):
                    r.text = part[1:-1]; r.font.name = "Consolas"; r.font.size = Pt(size - 1.5)
                    r.font.color.rgb = TEALD if r_i else WHITE
                    continue
                else:
                    r.text = part
                r.font.name = FONT; r.font.size = Pt(size)
                if r_i == 0:
                    r.font.bold = True; r.font.color.rgb = WHITE
                else:
                    r.font.color.rgb = INK
                    if bold_first and c_i == 0:
                        r.font.bold = True
            cell.fill.solid()
            cell.fill.fore_color.rgb = header_fill if r_i == 0 else (LIGHT if (zebra and r_i % 2 == 0) else WHITE)
    return shp


count = [0]


def slide(kicker, title, notes, dark=False):
    s = prs.slides.add_slide(BLANK)
    count[0] += 1
    n = count[0]
    if dark:
        rect(s, 0, 0, 13.333, 7.5, fill=NAVY2)
    box = rect(s, 0.55, 0.42, 0.36, 0.36, fill=TEALD if not dark else TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.15)
    text(s, 0.55, 0.42, 0.36, 0.36, f"{n:02d}", size=10, color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 1.05, 0.44, 11.7, 0.34, kicker.upper(), size=9.5, color=TEALD if not dark else TEAL, bold=True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.55, 0.86, 12.2, 0.95, title, size=23, color=NAVY if not dark else WHITE, bold=True, anchor=MSO_ANCHOR.TOP, line_sp=0.95)
    text(s, 0.55, 7.08, 8, 0.28, "NMIMS MPSTME  ·  Image & Video Processing", size=9, color=MUTED if not dark else PALE)
    text(s, 12.2, 7.08, 0.6, 0.28, f"{n:02d}", size=9, color=MUTED if not dark else PALE, align=PP_ALIGN.RIGHT)
    s.notes_slide.notes_text_frame.text = notes
    NOTES.append((n, title, notes))
    return s


def src(s, txt, y=6.78):
    text(s, 0.55, y, 12.2, 0.28, txt, size=8.5, color=MUTED, italic=True)


# =============================================================== 1 TITLE
s = prs.slides.add_slide(BLANK)
count[0] += 1
rect(s, 0, 0, 13.333, 7.5, fill=NAVY2)
for k, r_ in enumerate([5.2, 4.2, 3.2, 2.2, 1.2]):
    c = rect(s, 10.0 - r_ / 2 + 0.2, 3.6 - r_ / 2, r_, r_, fill=None, line=NAVY3 if k % 2 else RGBColor(0x1F, 0x3A, 0x6B), shape=MSO_SHAPE.OVAL, lw=1)
rect(s, 10.13, 3.53, 0.14, 0.14, fill=TEAL, shape=MSO_SHAPE.OVAL)
text(s, 0.8, 0.7, 9, 0.35, "NMIMS  MPSTME   ·   B.TECH COMPUTER ENGINEERING", size=10.5, color=TEAL, bold=True)
text(s, 0.8, 1.7, 9, 1.8, ["Digital Image", "Watermarking"], size=50, color=WHITE, bold=True, line_sp=0.9)
text(s, 0.8, 3.65, 8.6, 0.9, "A hybrid DWT–DCT–SVD approach — extended evaluation on 100 images, controls, technology comparison and a video feasibility test", size=15, color=PALE)
rect(s, 0.8, 4.75, 5.95, 0.42, fill=None, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.3)
text(s, 0.8, 4.75, 5.95, 0.42, "CLASSICAL SIGNAL PROCESSING  —  NO MACHINE LEARNING / DEEP LEARNING", size=9.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
text(s, 0.8, 5.95, 9, 0.9, ["**Krishang Vikram Singh**", "Roll No. B106  ·  Batch C1", "Subject: Image and Video Processing (IVP)  ·  First submission July 2026 · extended edition October 2026"], size=11.5, color=WHITE, spacing=2)
text(s, 0.55, 7.08, 8, 0.28, "NMIMS MPSTME  ·  Image & Video Processing", size=9, color=PALE)
s.notes_slide.notes_text_frame.text = ("Opening. This is the extended version of my watermarking project. The first submission tested a hybrid DWT-DCT-SVD scheme on two images. "
    "Since then I audited the code, rebuilt the evaluation on 100 unique images, added control experiments, compared the scheme with existing technology, and tried it on real video. "
    "I will show what held up, what did not, and what I corrected.")
NOTES.append((1, "Title", s.notes_slide.notes_text_frame.text))

# =============================================================== 2 MOTIVATION
s = slide("Introduction · motivation", "Copying an image takes a second — proving who owns it is the hard part",
          "Motivation: images are copied in one click. An invisible watermark carries an ownership signal inside the pixels. Three constraints must hold at once: imperceptible, robust, explainable. "
          "The project's third constraint is explainability: every step is a closed-form linear-algebra operation, no learned component. In this extended work I test whether the first two constraints hold at scale.")
card(s, 0.55, 2.0, 3.95, 2.55, "Imperceptible", "The watermark must not visibly change the image. Measured with PSNR and SSIM against the **unwatermarked original**.", size=14)
card(s, 4.69, 2.0, 3.95, 2.55, "Robust", "It must survive compression, noise, filtering and geometric change. Measured with NC / ZNCC / BER of the recovered logo.", size=14)
card(s, 8.83, 2.0, 3.95, 2.55, "Explainable", "DWT, DCT and SVD are closed-form and verifiable by hand — no black-box model. A deliberate scope choice, not a claim of superiority.", size=14)
rect(s, 0.55, 4.85, 12.23, 1.75, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, 0.85, 4.95, 11.6, 0.35, "WHAT THIS EXTENDED WORK ADDS", size=10, color=TEAL, bold=True)
bullets(s, 0.85, 5.3, 11.7, 1.3, ["**100 unique images** instead of two, with seeded attacks, status logging and 43 unit tests",
        "**Controls** that ask whether a high NC really proves the watermark is present",
        "**Technology comparison** (libraries and an existing watermarking package) and an exploratory **video** test"], size=12.5, color=WHITE, gap=3)

# =============================================================== 3 OBJECTIVES / SCOPE
s = slide("Objectives and scope", "Six objectives — three inherited from the first submission, three new",
          "Objectives one to three are from the first submission: implement, measure imperceptibility, measure robustness. The new ones are: characterise alpha, find and explain failure modes including false positives, and compare with existing technology and video. "
          "Scope boundaries matter: grey-scale only, no machine learning, non-blind scheme, video exploratory.")
table(s, 0.55, 1.95, 7.6, [
    ["1", "Implement the hybrid DWT–DCT–SVD embed/extract pipeline", "inherited"],
    ["2", "Quantify imperceptibility (PSNR, SSIM)", "inherited · now 100 images"],
    ["3", "Quantify robustness (NC, BER + ZNCC) under 14 attacks and a no-attack reference", "inherited · now 100 images"],
    ["4", "Characterise the α trade-off", "re-done, 4 attacks"],
    ["5", "Find failure modes — including false positives (controls)", "**new**"],
    ["6", "Compare with libraries and an existing package; test video", "**new**"],
], ["#", "Objective", "Status"], [0.4, 5.2, 2.0], size=11, row_h=0.52)
card(s, 8.45, 1.95, 4.33, 3.55, "Scope boundaries", [
    "8-bit **grey-scale** images (colour only via luminance)",
    "**No** machine / deep learning",
    "Scheme is **non-blind**: needs a key with the original singular values",
    "Video = **exploratory** (2 clips × 60 frames)",
    "Only operations that exist in the code are documented"], size=13)
rect(s, 0.55, 5.75, 12.23, 0.85, fill=TEALBG, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 0.8, 5.8, 11.8, 0.75, "**Honesty rule:** every number in this deck comes from a recorded experiment file; limits, unverified claims and corrections to the first submission are stated explicitly.", size=12.5, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

# =============================================================== 4 LITERATURE
s = slide("Existing technology · literature", "From LSB to hybrids — and an old warning about SVD schemes",
          "A short timeline. Spatial LSB marks by van Schyndel et al. 1994 are fragile. Cox et al. 1997 put a spread-spectrum watermark into significant spectral coefficients. "
          "Liu and Tan 2002 introduced SVD ownership watermarking; Ganic and Eskicioglu 2004 combined DWT and SVD. Later hybrids cascade DWT, DCT and SVD. "
          "The important warning: comments by Zhang and Li in 2005 and Ting in 2005 concern false positives and ambiguity attacks on this family. I verified the bibliographic details of each reference; "
          "I did not re-run those papers' experiments. The claim in my first report that hybrids always beat single transforms is a literature trend, not something I tested.")
xs = [0.55, 3.0, 5.45, 7.9, 10.35]
items = [("1994", "LSB marking", "van Schyndel et al. [10]", "high capacity, fragile"), ("1997", "Spread spectrum", "Cox et al. [1]", "significant spectral coefficients"),
         ("2002–04", "SVD & DWT-SVD", "Liu & Tan [2]; Ganic & Eskicioglu [3]", "singular-value embedding"),
         ("2010–22", "Hybrid DWT-DCT-SVD", "[7] [8] [9]", "reported gains; not re-tested here"), ("2005", "False-positive critique", "Zhang & Li [4]; Ting [5]", "key owner's mark can be 'found' anywhere")]
rect(s, 0.8, 2.52, 11.8, 0.05, fill=LINE)
for x, (yr, head, who, note) in zip(xs, items):
    red = head.startswith("False")
    rect(s, x + 0.9, 2.4, 0.28, 0.28, fill=RED if red else TEAL, shape=MSO_SHAPE.OVAL)
    text(s, x, 1.95, 2.2, 0.35, yr, size=12, color=RED if red else TEALD, bold=True, align=PP_ALIGN.CENTER)
    card(s, x, 2.95, 2.3, 2.1, head, [who, note], size=12, fill=REDBG if red else LIGHT, line=RED if red else LINE)
rect(s, 0.55, 5.35, 12.23, 1.3, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
bullets(s, 0.85, 5.45, 11.7, 1.2, ["Gap 1 — **identity binding**: SVD schemes can report the key owner's logo whatever was embedded (measured on slide 15)",
        "Gap 2 — **geometry**: global embedding breaks under cropping/rotation, the standard remedy being invariant or block-based domains [22]",
        "Gap 3 — **evaluation breadth**: published hybrids are typically shown on a few standard images; here 100"], size=11.5, color=WHITE, gap=2)
src(s, "Bibliographic details verified against publisher / secondary records; status flags in research/REFERENCES.md.", 6.72)

# =============================================================== 5 ARCHITECTURE (embedding)
s = slide("System architecture · embedding", "Embedding: add the logo's singular values to the host's — once, by the owner",
          "The embedding pipeline: host image, one-level Haar DWT, take the LL band, apply a 2-D DCT to the whole band, take the SVD, add alpha times the singular values of the logo, and invert the transforms. "
          "Two corrections to the first submission are visible here: the DCT is global, not block-wise. The result is an 8-bit image that looks like the original; the owner keeps a key.")
picture(s, V("diagram_embedding.png"), 0.55, 1.9, 12.2, 4.7)
src(s, "Correction vs first submission: the DCT is applied to the whole LL band (not 8×8 blocks). Alpha = 0.10 in the main experiment.", 6.72)

# =============================================================== 5b ARCHITECTURE (extraction)
s = slide("System architecture · verification", "Extraction needs the owner's key — and rebuilds the owner's logo from it",
          "Extraction repeats DWT, DCT and SVD on the received, possibly attacked image and subtracts the stored original singular values, divides by alpha and rebuilds the logo from the key's singular vectors. "
          "Second correction: the scheme is non-blind, not semi-blind, because the key contains the original singular values. And note the consequence we will measure later: because the logo's singular vectors are in the key, the extractor returns the key owner's logo whatever singular values it finds.")
picture(s, V("diagram_extraction.png"), 0.55, 1.9, 12.2, 4.7)
src(s, "Correction vs first submission: extraction is non-blind (needs S_orig of the unmarked host); the key holds U_w and V_w, i.e. the logo's identity.", 6.72)

# =============================================================== 6 MATHS
s = slide("Mathematical foundations", "One tunable parameter: add α times the logo's singular values",
          "The mathematics in five lines. Haar DWT gives the approximation band. A global orthonormal DCT-II moves it to a frequency domain. SVD factors the DCT matrix and the logo. "
          "The embedding rule adds alpha times the logo's singular values to the host's singular values. Extraction subtracts the stored original singular values and divides by alpha, then rebuilds the logo from the key's singular vectors. "
          "Important consequence: the identity of the logo lives in the key, not in the image.")
eqs = [("DWT (Haar, LL)", "LL = (a + b + c + d) / 2   for each 2×2 block"), ("DCT-II (global, orthonormal)", "F(u,v) = α(u)α(v) ΣΣ f(x,y) cos[(2x+1)uπ/2M] cos[(2y+1)vπ/2N]"),
       ("SVD", "A = U·diag(S)·Vᵀ        W = U_w·diag(S_w)·V_wᵀ"), ("Embedding", "S′ = S + α·S_w ,   A′ = U·diag(S′)·Vᵀ  →  IDCT → IDWT → clip"),
       ("Extraction (key: S_orig, U_w, V_w, α)", "S_w′ = (S_att − S_orig) / α ,   W′ = U_w·diag(S_w′)·V_wᵀ")]
for i, (h, e) in enumerate(eqs):
    y = 1.95 + i * 0.82
    rect(s, 0.55, y, 8.2, 0.7, fill=LIGHT, line=LINE, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
    text(s, 0.7, y + 0.03, 7.9, 0.28, h.upper(), size=9, color=TEALD, bold=True)
    text(s, 0.7, y + 0.28, 7.9, 0.4, e, size=14, color=NAVY, bold=False)
card(s, 9.05, 1.95, 3.73, 4.05, "What follows from the maths", [
    "**Non-blind:** needs S_orig of the unmarked host",
    "**Global DCT:** not tied to JPEG's 8×8 blocks — JPEG robustness comes from SVD",
    "**Identity in the key:** `U_w`, `V_w` rebuild *the owner's* logo from any singular values",
    f"Metric caution: a blank white image already scores NC = **{BASE['nc_constant_white_image']:.2f}** against this logo"], size=12.5)
src(s, "Equations (1)–(5) of the report; numerics unchanged from the first submission (re-run reproduces its table within 0.004 NC).", 6.72)

# =============================================================== 7 TRANSFORMATIONS
s = slide("Transformations implemented", "Twenty documented operations — only those that exist in the code",
          "Every operation in the code is documented in TRANSFORMATIONS_REFERENCE with its mathematics, parameters, limitations and complexity. The core is T2 to T5. "
          "T8 to T17 are the attacks used for testing. T20, video, is new and exploratory. Generic operations such as Sobel, CLAHE or morphology are not in this project and are not claimed.")
table(s, 0.55, 1.95, 12.23, [
    ["T1", "Grey conversion, centre-crop, Lanczos resize", "Pillow", "T11–T12", "Median 3×3, Gaussian blur σ=1.5"],
    ["T2", "1-level Haar DWT / inverse", "PyWavelets", "T13–T15", "Rotation 5°/15°, border cropping, 0.5× scaling"],
    ["T3", "2-D DCT-II / inverse (global)", "SciPy", "T16–T17", "Histogram equalisation, unsharp mask"],
    ["T4", "SVD of host sub-band and logo", "NumPy", "T18", "PSNR, SSIM, MSE"],
    ["T5", "Singular-value embedding and extraction", "NumPy", "T19", "NC, ZNCC (new), BER"],
    ["T6–T7", "Anti-aliased resize; procedural logo", "scikit-image, Pillow", "T20", "Frame I/O, per-frame embed, codec round trip (new)"],
    ["T8", "JPEG Q=90/50/20", "Pillow", "T9–T10", "Gaussian σ²=0.005/0.02; salt-and-pepper 2 %"],
], ["#", "Transformation", "Library", "#", "Transformation"], [0.7, 4.2, 1.6, 0.9, 4.2], size=10.5, row_h=0.5)
rect(s, 0.55, 5.9, 12.23, 0.75, fill=TEALBG, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 0.8, 5.93, 11.8, 0.7, "Pipeline order matters: DWT → DCT → SVD → embed → **round to 8 bit** → attack → DWT → DCT → SVD → extract. The first submission attacked the float image; the 8-bit protocol is primary here (difference < 0.001 mean NC except histogram equalisation).", size=11.5, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

# =============================================================== 8 DATASET
s = slide("Dataset design", f"{len(META)} unique images: {int((META.is_synthetic == 0).sum())} real, {int((META.is_synthetic == 1).sum())} labelled synthetic — rights and provenance recorded",
          "The dataset: 94 real images from open sample collections hosted on GitHub — OpenCV, scikit-image, torchvision, three.js textures and others — plus six synthetic controls. "
          "They are grey-standardised, square-cropped and resized to 512. Near duplicates were rejected by a hash, and a test asserts uniqueness. "
          "Important caveat: only raw GitHub and PyPI were reachable, so this is a convenience sample, not a benchmark, and rights for most photographs are not documented upstream — I record exactly what upstream says.")
picture(s, V("dataset_contact_sheet.png"), 0.55, 1.9, 7.0, 4.85)
cat = META.groupby("category").size().sort_values(ascending=False)
tile(s, 7.85, 1.95, 1.55, 1.15, "94", "real images", bigsize=26)
tile(s, 9.5, 1.95, 1.55, 1.15, "6", "synthetic controls", bigsize=26)
tile(s, 11.15, 1.95, 1.63, 1.15, "18", "categories", bigsize=26)
card(s, 7.85, 3.3, 4.93, 3.3, "Provenance, diversity, limits", [
    "Sources: OpenCV (+extra), scikit-image, darknet, torchvision, three.js, Pillow tests, …",
    f"Brightness: {', '.join(f'{k} {v}' for k, v in META.brightness_class.value_counts().items())}; contrast: {', '.join(f'{k} {v}' for k, v in META.contrast_class.value_counts().items())}",
    "SHA-256, URL, licence and rights note per image in `metadata.csv`",
    "Convenience sample — **not** a benchmark; colour removed; 1–4 images in several categories"], size=12)

# =============================================================== 9 EVALUATION DESIGN
s = slide("Quantitative evaluation methodology", "Seven experiments, 14 attacks plus a reference, four metrics — all seeded and logged",
          "E1 imperceptibility, E2 robustness with fifteen conditions — fourteen attacks plus the unattacked reference, which the first report counted as an attack — under two protocols, "
          "E3 an alpha sweep with four attacks, E4 false-positive controls. S1 and S2 are the technology comparisons and V1 is the video test. Metrics: PSNR and SSIM for imperceptibility; NC, ZNCC and BER for the recovered logo. "
          "The 0.75 threshold is the project's convention; I could not find a primary literature source for it, and I measure its false-alarm rate in E4.")
table(s, 0.55, 1.95, 7.9, [
    ["E1", "embed α = 0.10, save 8-bit PNG, verify", "100 images"],
    ["E2", "15 conditions × 2 protocols", "3000 trials"],
    ["E3", "α sweep × 4 attacks", "4000 rows"],
    ["E4", "false-positive controls (a–d)", "1600 rows"],
    ["S1/S2", "scheme / component benchmarks", "30 imgs / 10 ops"],
    ["V1", "video: 2 clips × 60 frames × 6 codecs", "720 frames"],
], ["ID", "What", "Size"], [0.8, 5.0, 2.0], size=11, row_h=0.46)
card(s, 8.75, 1.95, 4.03, 3.2, "Metrics", [
    "**PSNR / SSIM** vs the unwatermarked original",
    "**NC**, **ZNCC** (mean-centred), **BER** vs the logo",
    "Recognised if NC ≥ 0.75 (project convention)",
    "95 % bootstrap CIs; Wilson intervals for pass rates"], size=13)
rect(s, 0.55, 5.35, 12.23, 1.3, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.06)
text(s, 0.85, 5.42, 11.7, 1.2, [f"**{RS['images_ok']}/{RS['images_requested']}** images · **{RS['attack_trials_ok']}/{RS['attack_trials_total']}** attack trials · **{RS['alpha_sweep_rows_ok']}/{RS['alpha_sweep_rows_total']}** sweep rows — **0 failures**",
        "Two full runs gave **identical metrics** (max |Δ| = 0.0) · 43 tests pass · ≈ 80 s on 4 cores"], size=13, color=WHITE, spacing=4)

# =============================================================== 10 IMPERCEPTIBILITY
s = slide("Results · imperceptibility", "Invisible on typical images — but PSNR alone hides some failures",
          f"Across 100 images the mean PSNR is {E1['psnr_all']['mean']:.1f} dB and mean SSIM {E1['ssim_all']['mean']:.3f}. Notice that PSNR is almost the same for most images, about 35.6 dB, "
          "because the added energy is set by the logo's singular values, not by the host. SSIM exposes exceptions: the near-black Milky Way image has PSNR 35.9 but SSIM 0.58, and the checkerboard 0.86. "
          "So PSNR alone would have hidden these failures.")
tile(s, 0.55, 1.95, 2.3, 1.2, f"{E1['psnr_all']['mean']:.1f} dB", "mean PSNR (min 35.5)", bigsize=28)
tile(s, 0.55, 3.3, 2.3, 1.2, f"{E1['ssim_all']['mean']:.3f}", f"mean SSIM · {E1['n_ssim_below_0.98']} images < 0.98", bigsize=28)
tile(s, 0.55, 4.65, 2.3, 1.2, "0", "images below 30 dB", bigsize=28)
picture(s, V("fig_imperceptibility_hist.png"), 3.05, 1.9, 5.3, 2.2)
picture(s, V("fig_representative_comparison.png"), 8.6, 1.85, 4.2, 4.9)
bullets(s, 3.05, 4.25, 5.4, 2.4, ["PSNR ≈ 35.6 dB for **82 of 100** images — set by the logo, not the host",
        "**Near-black sky** (`image_055`): PSNR 35.9 dB but **SSIM 0.58** — clipping dominates",
        "Checkerboard: SSIM 0.86; synthetic images reach 43.8 dB"], size=11.5, gap=4)
src(s, "Right: original · watermarked · |difference|×10 for typical images and two failure cases (E1, 100 images).", 6.78)

# =============================================================== 11 ROBUSTNESS
s = slide("Results · robustness under attack", f"{pc(E2['pass_rate'], 1)} of trials recovered — strong for compression and filtering, weak for cropping and equalisation",
          f"Headline: {E2['n_pass']} of 1500 trials reach NC at least 0.75 when the unattacked reference is included, as in the first submission, and {E2['n_pass_attacks_only']} of {E2['n_trials_attacks_only']} counting only the 14 attacks. "
          "Very robust: JPEG, median filter, sharpening, scaling, salt and pepper, blur, mild noise. Moderate: five and fifteen degree rotation, strong noise. "
          "Weak: histogram equalisation, 44 percent, and border cropping, 1 percent. The first submission said moderate rotation and scaling had mean NC above 0.93; with 100 images that holds for scaling, 0.938, but not for 5 degree rotation, 0.880.")
picture(s, V("fig_nc_boxplot.png"), 0.55, 1.85, 7.4, 4.85)
cols = [("JPEG Q=50", "JPEG Q=50"), ("Median 3×3", "Median Filter 3x3"), ("Scaling 0.5×", "Scaling 0.5x"), ("Noise σ²=0.02", "Gaussian Noise (var=0.02)"),
        ("Rotation 15°", "Rotation 15 deg"), ("Hist. equalisation", "Histogram Equalization"), ("Border cropping", "Cropping 25%")]
rows = [[a, f"{ROB.loc[b].mean_nc:.3f}", f"{100 * ROB.loc[b].pass_rate:.0f} %"] for a, b in cols]
table(s, 8.2, 1.95, 4.58, rows, ["Attack", "mean NC", "NC ≥ 0.75"], [2.2, 1.1, 1.3], size=11, row_h=0.42)
rect(s, 8.2, 5.5, 4.58, 1.15, fill=TEALBG, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 8.35, 5.53, 4.3, 1.1, "First submission: \"rotation/scaling NC > 0.93\". At n = 100: scaling 0.938 ✓, **5° rotation 0.880 ✗**.", size=11, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)
src(s, f"1500 trials = 100 images × 15 conditions (14 attacks + reference); attacks only: {E2['n_pass_attacks_only']}/{E2['n_trials_attacks_only']} = {pc(E2['pass_rate_attacks_only'], 1)}.", 6.78)

# =============================================================== 12 WHICH IMAGES FAIL
s = slide("Results · which images fail?", "Robustness depends on the image: higher-contrast images lose more under noise and rotation",
          "Across the 100 images, NC after noise and rotation is negatively correlated with image contrast — Spearman rho about minus 0.65 for strong Gaussian noise, minus 0.68 for salt and pepper and minus 0.55 for 15 degree rotation. "
          "Histogram equalisation instead correlates with mean brightness. These are associations, with uncorrected p-values and n of 100, not demonstrated causes. "
          "By category, the two document-text images are the weakest, portraits and documents are most sensitive to equalisation; several categories have only one to four images so I treat them as descriptive.")
picture(s, V("fig_nc_vs_contrast.png"), 0.55, 1.9, 6.4, 2.25)
picture(s, V("fig_category_heatmap.png"), 7.2, 1.85, 5.6, 4.9)
bullets(s, 0.55, 4.35, 6.4, 2.4, [f"Noise σ²=0.02: ρ(contrast) = **{COR.loc['Gaussian Noise (var=0.02)'].rho_contrast:.2f}**; salt-and-pepper **{COR.loc['Salt & Pepper (2%)'].rho_contrast:.2f}**; 15° rotation **{COR.loc['Rotation 15 deg'].rho_contrast:.2f}**",
        f"Hist. equalisation tracks mean brightness (ρ = {COR.loc['Histogram Equalization'].rho_mean_intensity:.2f}), not contrast",
        "Weakest categories: **document text** (noise 0.53) and **portraits** under equalisation (0.40)",
        "n = 100, uncorrected p-values — associations, not causes"], size=11.5, gap=4)

# =============================================================== 13 ALPHA
s = slide("Results · embedding strength", "The α trade-off is real for noise and rotation — but flat for JPEG, so the earlier figure misled",
          "This slide corrects the first submission. Its trade-off figure used two y-axes with the NC axis zoomed to 0.93 to 0.99, and its JPEG simulation truncated floats, which lowered NC only at alpha 0.02. "
          "I re-created its conditions exactly and got the same numbers, then showed that rounding removes the dip. With 100 images and four attacks the picture is clear: after JPEG the NC is flat, "
          "but under noise, blur and rotation robustness grows with alpha while PSNR falls. For Gaussian noise the pass rate goes from 9 percent at alpha 0.02 to 85 percent at 0.10 and 95 percent at 0.40, while PSNR falls below 30 dB from alpha 0.20.")
picture(s, V("fig_alpha_tradeoff.png"), 0.55, 1.9, 12.2, 3.6)
g = "Gaussian Noise (var=0.02)"
tile(s, 0.55, 5.65, 2.9, 1.05, f"{pc(sw(g, 0.02, 'pass_rate'))} → {pc(sw(g, 0.10, 'pass_rate'))}", "noise pass rate, α 0.02 → 0.10", bigsize=22)
tile(s, 3.6, 5.65, 2.9, 1.05, f"{sw('No Attack', 0.10, 'psnr'):.1f} → {sw('No Attack', 0.20, 'psnr'):.1f} dB", "PSNR, α 0.10 → 0.20", bigsize=22)
tile(s, 6.65, 5.65, 2.9, 1.05, f"{sw('JPEG Q=50', 0.02, 'nc_mean'):.2f} → {sw('JPEG Q=50', 0.40, 'nc_mean'):.2f}", "mean NC after JPEG Q=50, α 0.02 → 0.40", bigsize=22)
rect(s, 9.7, 5.65, 3.08, 1.05, fill=REDBG, line=RED, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 9.82, 5.68, 2.85, 1.0, "Old figure: axis 0.93–0.99 + float truncation before JPEG → **artefact** (`exp_jpeg_truncation.py`)", size=10.5, color=INK, anchor=MSO_ANCHOR.MIDDLE)

# =============================================================== 14 FALSE POSITIVE
s = slide("Results · false-positive controls", f"A different logo decoded with the owner's key still scores NC {CC.mean_nc:.2f} — NC does not prove ownership",
          "This is the most important finding. Control C: embed a different logo in the same host and decode with the key of logo A. The extractor returns the key owner's logo, NC 0.975 against 0.977 for the genuine pair. "
          "That is the false-positive problem described for SVD schemes in the literature, now measured for this implementation. Controls A and B: unrelated images decoded with a wrong key score 0.37 to 0.40 on average, "
          "and about 10 to 13 percent exceed the 0.75 threshold. A claim of ownership therefore needs an additional safeguard such as an image-bound signature; I did not implement one. "
          "I also tried embedding the principal component instead; that made things worse, the unwatermarked image still decoded at 0.98.")
picture(s, V("fig_false_positive.png"), 0.55, 1.9, 6.3, 3.4)
picture(s, V("fig_controls_hist.png"), 6.95, 1.9, 5.85, 3.15)
tile(s, 0.55, 5.5, 2.95, 1.15, f"{CC.mean_nc:.3f}", "wrong logo, owner's key (genuine 0.977)", bigc=RGBColor(0xFF, 0x8A, 0x80), bigsize=26)
tile(s, 3.65, 5.5, 2.95, 1.15, f"{pc(CA.frac_ge_0_75, 1)} / {pc(CB.frac_ge_0_75, 1)}", "unrelated-image controls above NC 0.75", bigsize=22)
tile(s, 6.75, 5.5, 2.95, 1.15, f"{E4['wrong_image_p95']:.2f}", "95th percentile NC of unrelated images", bigsize=26)
rect(s, 9.85, 5.5, 2.93, 1.15, fill=TEALBG, line=TEAL, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.08)
text(s, 9.97, 5.52, 2.7, 1.1, "Needed: an **image-bound signature** — not implemented; the principal-component variant was **worse** (NC 0.98 on an unmarked image)", size=10, color=NAVY, anchor=MSO_ANCHOR.MIDDLE)

# =============================================================== 15 FAILURE CASES
s = slide("Results · failure case analysis", "Four failure modes — explained from the data, not hidden",
          "Border cropping: mean NC 0.02; with 43.75 percent of the area zeroed every singular value changes. The first report called this 25 percent cropping; 25 percent of each dimension is 43.75 percent of the area. "
          "Histogram equalisation: 44 percent pass; it varies between images, 0.10 for the cameraman and 0.66 for the astronaut in the first submission. "
          "Near-constant images: SSIM collapses though PSNR is fine. Rotation: 72 percent pass at 15 degrees; there is no re-registration.")
card(s, 0.55, 1.95, 2.95, 3.9, "Border cropping", [f"mean NC **{ROB.loc['Cropping 25%'].mean_nc:.3f}**, pass **{pc(ROB.loc['Cropping 25%'].pass_rate)}**", "43.75 % of the *area* zeroed (25 % per dimension)", "all singular values change; logo extracts as black (BER 0.14 ≈ white fraction)"], fill=REDBG, line=RED, tcolor=RED, size=12.5)
card(s, 3.65, 1.95, 2.95, 3.9, "Histogram equalisation", [f"mean NC **{ROB.loc['Histogram Equalization'].mean_nc:.3f}**, pass **{pc(ROB.loc['Histogram Equalization'].pass_rate)}**", "remaps the whole intensity distribution", "varies by image (first report: 0.10 vs 0.66)", "tracks mean brightness (ρ −0.31)"], fill=REDBG, line=RED, tcolor=RED, size=12.5)
card(s, 6.75, 1.95, 2.95, 3.9, "Rotation > 5°", [f"5°: **{pc(ROB.loc['Rotation 5 deg'].pass_rate)}**, 15°: **{pc(ROB.loc['Rotation 15 deg'].pass_rate)}**", "no re-registration attempted", "consistent with RST theory [22]", "block-based schemes break too (blind-watermark BER 0.21)"], fill=REDBG, line=RED, tcolor=RED, size=12.5)
card(s, 9.85, 1.95, 2.93, 3.9, "Near-constant images", ["near-black sky: PSNR **35.9 dB**, SSIM **0.58**", "checkerboard: lowest clean NC (0.79)", "PSNR alone would have hidden both"], fill=REDBG, line=RED, tcolor=RED, size=12.5)
rect(s, 0.55, 6.05, 12.23, 0.6, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.1)
text(s, 0.8, 6.07, 11.8, 0.56, f"By contrast: JPEG Q ≥ 20, median 3×3, sharpening, 0.5× scaling, blur, salt-and-pepper and mild noise all reach ≥ 95 % pass rate.", size=12, color=WHITE, anchor=MSO_ANCHOR.MIDDLE)

# =============================================================== 16 COMPONENT COMPARISON
s = slide("Comparison with existing technologies · components", "Same operation, different libraries: numerically equal, speed rarely matters here",
          "Component benchmark: the same operation on the same input in different libraries, one thread, in memory, warm-up and repeats. Results agree numerically: DCT and SVD to ten to the minus sixteen, "
          "Pillow and OpenCV JPEG bit-identical. OpenCV is much faster for the attack operations — about 340 times for the 3 by 3 median — partly because it works on uint8. But the whole 100-image experiment takes about 80 seconds, "
          "so changing libraries is not justified on speed. ImageMagick as a command line tool costs about 49 milliseconds per call, dominated by process start-up.")
picture(s, V("fig_tech_speedup.png"), 0.55, 1.95, 6.9, 4.2)
card(s, 7.65, 1.95, 5.13, 4.7, "What the measurements say", [
    "DCT / SVD / Haar agree to ≤ 1.5e-16 (DWT 1e-13); Pillow ≡ OpenCV JPEG",
    "SVD dominates embedding (≈ 23 of 37 ms); `scipy.linalg.svd` is **1.19×** NumPy's",
    "OpenCV: median **≈ 340×**, blur 4.4×, rotation 5.1×, rescale 38×, equalise 19× (partly `uint8` vs `float64`)",
    "PyWavelets Haar is **11× slower** than four lines of NumPy — still the right general tool",
    "ImageMagick CLI ≈ **49 ms** per JPEG (process + disk)",
    "Whole experiment ≈ **80 s** on 4 cores → no performance case for switching"], size=11.5)
src(s, "Documentation-based capability matrix (not benchmarked) is in the report, Table 3.1.", 6.78)

# =============================================================== 17 SCHEME COMPARISON
s = slide("Comparison with existing technologies · schemes", "Beats LSB and block-DCT at equal PSNR — but not an existing blind block-based hybrid",
          "Scheme-level comparison with the same 4095-bit payload and the same attacks on 30 images. The block-DCT baseline was tuned to match the project's PSNR. LSB is imperceptible but destroyed by any processing. "
          "The open-source blind-watermark package, which is also DWT-DCT-SVD but block-based and blind, has lower bit-error rates under most attacks at higher PSNR, 42 against 36 dB. "
          "But it is 13 times slower to embed, 18 times slower to extract, and worse under rotation and strong JPEG. Caveat: its YUV design and my channel-wise attacks may favour it. "
          "So the project is a simple, explainable, key-based alternative, not a better algorithm.")
picture(s, V("fig_tech_scheme_ber.png"), 0.55, 1.85, 7.9, 3.4)
rows = [["PSNR (dB)"] + [f"{psnr_s(x):.1f}" for x in (PROJ, BWM, BDCT, LSB)],
        ["JPEG Q=50 BER"] + [f"{ber(x, 'JPEG Q=50'):.3f}" for x in (PROJ, BWM, BDCT, LSB)],
        ["Rotation 5° BER"] + [f"{ber(x, 'Rotation 5 deg'):.3f}" for x in (PROJ, BWM, BDCT, LSB)],
        ["Cropping BER"] + [f"{ber(x, 'Cropping 25%'):.3f}" for x in (PROJ, BWM, BDCT, LSB)],
        ["embed / extract ms"] + [f"{tms(x, 't_embed_ms'):.0f} / {tms(x, 't_extract_ms'):.0f}" for x in (PROJ, BWM, BDCT, LSB)],
        ["blind?"] + ["**no** (key)", "yes", "yes", "yes"]]
table(s, 0.55, 5.3, 7.9, rows, ["", "Project", "blind-watermark", "block DCT", "LSB"], [2.2, 1.5, 1.8, 1.4, 1.2], size=10, row_h=0.235)
card(s, 8.7, 1.95, 4.08, 4.7, "Reading the comparison", [
    "Project ≫ LSB and block-DCT under JPEG, scaling, rotation",
    f"`blind-watermark`: lower BER under JPEG Q=50 ({ber(BWM, 'JPEG Q=50'):.3f} vs {ber(PROJ, 'JPEG Q=50'):.3f}), noise, blur, cropping; **higher PSNR** ({psnr_s(BWM):.1f} dB)",
    f"…but **{tms(BWM, 't_embed_ms') / tms(PROJ, 't_embed_ms'):.0f}× slower** and worse under rotation ({ber(BWM, 'Rotation 5 deg'):.3f} vs {ber(PROJ, 'Rotation 5 deg'):.3f}) and JPEG Q=20",
    "Its colour/YUV design and channel-wise attacks may favour it",
    "**No claim of superiority** — contribution = explainable reference implementation + measured weaknesses"], size=12)

# =============================================================== 18 VIDEO
s = slide("Exploratory video test", "Per-frame embedding survives H.264 and MJPEG re-encoding — feasibility, not a video benchmark",
          "Because the project's future-scope slide mentioned video, I tested it: 60 frames from each of two real clips, read with OpenCV, watermarked frame by frame, each with its own key, re-encoded with FFmpeg as lossless FFV1, H.264 at three quality levels, and MJPEG at two. "
          "All 720 decoded frames gave NC at least 0.972. So lossy codecs are not the difficulty. The caveats are important: keys are per frame, about 2 megabytes each; temporal attacks such as frame dropping or re-timing were not tested; "
          "and a naive misaligned frame-drop test still gave NC 0.97, presumably because neighbouring frames are similar — another sign that NC does not bind a mark to a frame.")
v = VID.groupby("codec", sort=False).agg(nc=("nc_mean", "mean"), ncmin=("nc_min", "min"), psnr=("psnr_wm_vs_orig_db", "mean"), kb=("size_wm_kb", "mean")).reset_index()
rows = [[x.codec, f"{x.psnr:.1f}", f"{x.nc:.3f}", f"{x.ncmin:.3f}", f"{x.kb:.0f}"] for x in v.itertuples()]
table(s, 0.55, 1.95, 7.6, rows, ["Codec setting", "PSNR (dB)", "mean NC", "min NC", "size (kB)"], [2.6, 1.2, 1.2, 1.2, 1.2], size=11, row_h=0.46)
tile(s, 8.45, 1.95, 2.1, 1.2, "720", "decoded frames, 0 failed", bigsize=28)
tile(s, 10.68, 1.95, 2.1, 1.2, f"≥ {VID.nc_min.min():.2f}", "NC in every frame", bigsize=28)
card(s, 8.45, 3.35, 4.33, 2.7, "Caveats", ["2 clips × 60 frames — small", "one **key per frame** (≈ 2 MB)", "no frame dropping, re-timing, colour or hardware codecs", "naive frame-drop still NC 0.97 (similar neighbours)"], size=12, fill=REDBG, line=RED, tcolor=RED)
text(s, 0.55, 5.35, 7.6, 1.0, f"Embedding ≈ 32–41 ms per 512² frame on one core (≈ 24–31 frames/s) · clips: `vtest.avi` (768×576), `Megamind.avi` (720×528) from the OpenCV sample data", size=11, color=MUTED)

# =============================================================== 19 CORRECTIONS
s = slide("Claims versus evidence", "What held up and what did not: ten claims of the first submission checked at n = 100",
          "A direct comparison of the first submission's claims with the 100-image evidence. The headline numbers are reproducible and mostly hold. "
          "What changed: moderate rotation is weaker than claimed, the alpha trade-off figure was an artefact for JPEG, the block-DCT and semi-blind labels were wrong, "
          "the suite is fourteen attacks plus a reference, and — most importantly — extraction does not prove ownership.")
rows = [
    ["PSNR 35.6 dB / SSIM 0.996 (2 images)", f"{E1['psnr_all']['mean']:.1f} dB / {E1['ssim_all']['mean']:.3f}; 3 images < 0.98", "✓ with exceptions"],
    ["25 / 30 trials ≥ 0.75", f"{E2['n_pass']} / 1500 ({pc(E2['pass_rate'], 1)}); attacks only {pc(E2['pass_rate_attacks_only'], 1)}", "✓"],
    ["JPEG: NC > 0.97", f"{ROB.loc['JPEG Q=90'].mean_nc:.3f} / {ROB.loc['JPEG Q=50'].mean_nc:.3f} / {ROB.loc['JPEG Q=20'].mean_nc:.3f} (Q 90/50/20)", "✓ (Q ≥ 50)"],
    ["Rotation/scaling NC > 0.93", f"scaling {ROB.loc['Scaling 0.5x'].mean_nc:.3f}; rotation 5° {ROB.loc['Rotation 5 deg'].mean_nc:.3f}", "scaling ✓ · rotation ✗"],
    ["Cropping fails; hist-eq 0.10 / 0.66", f"pass {pc(ROB.loc['Cropping 25%'].pass_rate)}; hist-eq mean {ROB.loc['Histogram Equalization'].mean_nc:.3f} (44 % pass)", "✓ (hist-eq varies)"],
    ["Robustness rises with α (JPEG)", f"flat {sw('JPEG Q=50', 0.02, 'nc_mean'):.3f} → {sw('JPEG Q=50', 0.40, 'nc_mean'):.3f}", "✗ artefact"],
    ["\"Block DCT\", \"semi-blind\"", "global DCT, non-blind", "corrected"],
    ["\"Fifteen attacks\"", "14 attacks + the unattacked reference", "corrected"],
    ["Extraction proves ownership", f"wrong logo decoded at NC {CC.mean_nc:.3f}", "✗ **not supported**"],
    ["Hybrids beat single transforms", "not tested; blind-watermark beats the project", "unverified"],
]
table(s, 0.55, 1.9, 12.23, rows, ["First submission", "This evaluation (100 images)", "Verdict"], [4.4, 5.6, 2.2], size=12, row_h=0.45)

# =============================================================== 20 IMPLEMENTATION
s = slide("Implementation · quality · reproducibility", "Tested, seeded and logged: 43 tests, zero failures, identical re-runs",
          "Implementation: the embed and extract numerics are unchanged, so the first submission's numbers remain reproducible — I re-ran its experiment and got the same table within 0.004. "
          "Added: input validation, a safe image loader, a dataset builder, a multiprocessing batch runner that records status and time per operation and re-opens every output, seeded attacks, and 43 tests. "
          "Bug found during development, for transparency: a debug flag overwrote the metadata file; fixed, with a regression test.")
card(s, 0.55, 1.95, 4.0, 3.0, "Engineering", ["numerics of `embed/extract_watermark` unchanged", "validation: 2-D, finite, size, α > 0", "`image_io`: alpha, 16-bit, corrupt/empty files", "config-driven batch runner, 4 workers"], size=12.5)
card(s, 4.7, 1.95, 4.0, 3.0, "Verification", ["**43 tests** pass (functions, edge cases, loaders, completeness)", "legacy re-run reproduces the first report (≤ 0.004 NC)", "two full runs: identical metrics", "all 100 output PNGs re-opened and validated"], size=12.5)
card(s, 8.85, 1.95, 3.93, 3.0, "Traceability", ["originals preserved in `original_backup/`", "CHANGELOG lists every change and why", "negative results saved: principal-component variant; JPEG truncation", "PROJECT_AUDIT: 15 findings, F1–F15"], size=12.5)
rect(s, 0.55, 5.2, 12.23, 1.45, fill=NAVY, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.05)
text(s, 0.85, 5.28, 11.7, 0.3, "REPRODUCE", size=10, color=TEAL, bold=True)
text(s, 0.85, 5.58, 11.7, 1.0, ["`python build_dataset.py  →  run_batch_experiment.py  →  analysis.py  →  benchmark_tech.py  →  video_demo.py`", "`python -m pytest tests -q`   ·   Python 3.13, NumPy 2.5, SciPy 1.18, scikit-image 0.26, Pillow 12.3, OpenCV 5.0, FFmpeg 6.1"], size=11.5, color=WHITE, spacing=3)

# =============================================================== 21 LIMITATIONS & FUTURE
s = slide("Limitations · future scope", "Known limits — and the next steps the evidence supports",
          "Limitations: false positives, non-blind key, cropping and rotation, histogram equalisation, grey-scale only, a lossy clean extraction of about 1 percent BER, and a dataset that is a convenience sample. "
          "Future work is ordered by evidence: first bind the logo to the image; then block-based redundant embedding for cropping; then geometric re-synchronisation; colour via luminance; and a proper video design with shared or index-aware keys.")
card(s, 0.55, 1.95, 5.95, 4.7, "Limitations (measured)", [
    "**False-positive problem** — not fixed",
    "**Non-blind**: ≈ 2.1 MB key per image",
    "**Cropping** 1 %, **rotation 15°** 72 %, **hist-eq** 44 %",
    "Grey-scale only; clean BER ≈ 0.9 % (logo resampling)",
    "Dataset: convenience sample, rights mostly undocumented",
    "Only one existing package compared; no GPU; ±5–10 % timing noise",
    "Video: exploratory, per-frame keys, no temporal attacks"], size=13, fill=REDBG, line=RED, tcolor=RED)
card(s, 6.7, 1.95, 6.08, 4.7, "Future scope (in priority order)", [
    "**1 · Image-bound signature** to remove false positives — verify with the E4 controls",
    "**2 · Block-based redundant embedding** for cropping; use `blind-watermark` as the new baseline",
    "**3 · Geometric re-synchronisation** (log-polar / template) [22]",
    "**4 · Colour:** embed in the Y channel of YCbCr",
    "**5 · Video:** index-aware keys, temporal redundancy, frame-drop and re-timing tests; real-time port (≈ 27 fps single-core as is)",
    "**6 ·** Learned watermarkers [29][30] as a comparison, if the no-ML rule is relaxed"], size=13)

# =============================================================== 22 CONCLUSION
s = slide("Conclusion", "A transparent scheme that survives everyday processing — but NC alone is not proof of ownership",
          "Conclusion. The scheme is transparent and survives re-compression and filtering well: 85 percent of all trials, 84 percent of attack trials. "
          "Imperceptibility is good on typical images with exceptions. The extension changed the story in four ways: the alpha trade-off exists only for harsher attacks, rotation and equalisation are weaker than claimed, "
          "a high NC does not prove ownership, and an existing package beats the project on most attacks. What I would do next is bind the logo to the image. Thank you; questions welcome.", dark=True)
tile(s, 0.55, 1.95, 3.9, 1.5, f"{E1['psnr_all']['mean']:.1f} dB / {E1['ssim_all']['mean']:.3f}", "mean PSNR / SSIM over 100 images", fill=NAVY3, bigsize=28)
tile(s, 4.7, 1.95, 3.9, 1.5, f"{pc(E2['pass_rate'], 1)}", "of 1500 trials with NC ≥ 0.75", fill=NAVY3, bigsize=28)
tile(s, 8.85, 1.95, 3.93, 1.5, f"NC {CC.mean_nc:.2f}", "for a DIFFERENT logo with the owner's key", fill=NAVY3, bigc=RGBColor(0xFF, 0x8A, 0x80), bigsize=28)
text(s, 0.55, 3.75, 12.2, 0.3, "KEY TAKEAWAYS", size=10, color=TEAL, bold=True)
bullets(s, 0.55, 4.1, 12.2, 2.7, ["**Works:** JPEG Q ≥ 20, median, sharpening, scaling, blur, mild noise — ≥ 95 % pass rate; per-frame video re-encoding too",
        "**Does not:** border cropping (1 %), histogram equalisation (44 %), rotation 15° (72 %); first report's α-trade-off figure was an artefact for JPEG",
        "**Cannot claim:** ownership from NC alone (false-positive problem) or superiority over existing technology",
        "**Next:** bind the logo to the image, block-based redundancy, geometric re-synchronisation"], size=15, color=WHITE, gap=9)

# =============================================================== 23 REFERENCES
s = slide("References", "Verified sources (IEEE style) — status flags in research/REFERENCES.md",
          "References. All bibliographic details were checked against publisher or secondary records; entries that could not be verified were not cited. The full list with verification status is in research/REFERENCES.md. Thank you.")
refs = [
    "[1] I. J. Cox et al., “Secure spread spectrum watermarking for multimedia,” IEEE Trans. Image Process., 6(12):1673–1687, 1997.",
    "[2] R. Liu and T. Tan, “An SVD-based watermarking scheme for protecting rightful ownership,” IEEE Trans. Multimedia, 4:121–128, 2002.",
    "[3] E. Ganic and A. M. Eskicioglu, “Robust DWT-SVD domain image watermarking: embedding data in all frequencies,” Proc. MM&Sec ’04, ACM, pp. 166–174, 2004.",
    "[4] X.-P. Zhang and K. Li, “Comments on ‘An SVD-based watermarking scheme…’,” IEEE Trans. Multimedia, 7:593–594, 2005.",
    "[5] G. C.-W. Ting, “Ambiguity attacks on the Ganic–Eskicioglu DWT-SVD scheme,” Proc. ICISC, pp. 378–388, 2005.",
    "[7] C.-C. Lai and C.-C. Tsai, “Digital image watermarking using DWT and SVD,” IEEE Trans. Instrum. Meas., 59(11):3060–3063, 2010.",
    "[8] S. Bekkouch and K. M. Faraoun, “…combined DCT-DWT-SVD transforms,” J. Inf. Process. Syst., 11(3):406–420, 2015.",
    "[9] M. Begum, J. Ferdush, M. S. Uddin, “A hybrid robust watermarking system based on DCT, DWT and SVD,” J. King Saud Univ. CIS, 34(8):5856–5867, 2022.",
    "[10] R. G. van Schyndel et al., “A digital watermark,” Proc. ICIP, vol. 2, pp. 86–90, 1994.   [11] E. Koch and J. Zhao, Proc. IEEE NSIP Workshop, pp. 452–455, 1995.",
    "[12] S. G. Mallat, IEEE Trans. PAMI, 11(7):674–693, 1989.   [13] N. Ahmed, T. Natarajan, K. R. Rao, IEEE Trans. Comput., C-23(1):90–93, 1974.",
    "[16] Z. Wang et al., “Image quality assessment: from error visibility to structural similarity,” IEEE Trans. Image Process., 13(4):600–612, 2004.",
    "[17] F. Hartung and M. Kutter, “Multimedia watermarking techniques,” Proc. IEEE, 87(7):1079–1107, 1999.",
    "[22] C.-Y. Lin et al., “Rotation, scale, and translation resilient watermarking for images,” IEEE Trans. Image Process., 10(5):767–782, 2001.",
    "[23] F. Hartung and B. Girod, “Watermarking of uncompressed and compressed video,” Signal Process., 66(3):283–301, 1998.",
    "[28] `blind-watermark` 0.4.4, github.com/guofei9987/blind_watermark   ·   [29] J. Zhu et al., HiDDeN, ECCV 2018   ·   [30] M. Tancik et al., StegaStamp, CVPR 2020",
    "Software: NumPy [19], SciPy [20], scikit-image [21], PyWavelets [18], Pillow, OpenCV, FFmpeg, ImageMagick.   Dataset sources and attributions: data/README.md.",
]
text(s, 0.55, 1.85, 12.2, 4.95, refs, size=11.5, color=INK, spacing=5)

prs.save(os.path.join(OUT, "Final_Presentation.pptx"))

# ------------------------------------------------------------------ speaker notes file
with open(os.path.join(OUT, "SPEAKER_NOTES.md"), "w", encoding="utf-8") as f:
    f.write("# Speaker notes — Final_Presentation.pptx\n\nTarget length: ≈ 15–18 minutes for 24 slides (≈ 40–45 s per slide; slides 12, 14, 15 and 18 deserve more time).\n\n")
    for n, title, notes in NOTES:
        f.write(f"## Slide {n} — {title}\n\n{notes}\n\n")
print("slides:", count[0])
