"""Render the synthetic sample knowledge base into data/knowledge_base/.

Usage: python scripts/build_sample_kb.py
"""
from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parent))
from kb_content import DOCS  # noqa: E402

OUT = Path(__file__).resolve().parent.parent / "data" / "knowledge_base"


def write_md(path: Path, title: str, secs) -> None:
    lines = [f"# {title}", ""]
    for h, p in secs:
        lines += [f"## {h}", "", p, ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_txt(path: Path, title: str, secs) -> None:
    lines = [title.upper(), "=" * len(title), ""]
    for h, p in secs:
        lines += [h, "-" * len(h), p, ""]
    path.write_text("\n".join(lines), encoding="utf-8")


def write_docx(path: Path, title: str, secs) -> None:
    from docx import Document

    d = Document()
    d.add_heading(title, 0)
    for h, p in secs:
        d.add_heading(h, 1)
        d.add_paragraph(p)
    d.save(path)


def write_pdf(path: Path, title: str, secs) -> None:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer

    ss = getSampleStyleSheet()
    # Small fixed header/footer so the cleaner's repeated-line removal is exercised.
    def deco(c, doc):
        c.setFont("Helvetica", 8)
        c.drawString(50, 810, "RAG Study Guide (synthetic sample corpus)")
        c.drawString(50, 25, f"Page {doc.page}")

    story = [Paragraph(title, ss["Title"]), Spacer(1, 12)]
    for i, (h, p) in enumerate(secs):
        story += [Paragraph(h, ss["Heading2"]), Paragraph(p, ss["BodyText"]), Spacer(1, 10)]
        if i == 2:  # force a second page so page metadata is meaningful
            story.append(PageBreak())
    SimpleDocTemplate(str(path), pagesize=A4, title=title, author="Project authors (synthetic)").build(
        story, onFirstPage=deco, onLaterPages=deco)


WRITERS = {"md": write_md, "txt": write_txt, "docx": write_docx, "pdf": write_pdf}

if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    for doc_id, (title, fmt, secs) in DOCS.items():
        p = OUT / f"{doc_id}.{fmt}"
        WRITERS[fmt](p, title, secs)
        print("wrote", p.name)
