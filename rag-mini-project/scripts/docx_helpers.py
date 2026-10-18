"""python-docx helpers for consistent, professional deliverables."""
from __future__ import annotations

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

FONT = "Calibri"


def new_doc(title_for_footer: str, base_pt: float = 10.5, margin_cm: float = 2.0) -> Document:
    d = Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(margin_cm)
    sec.top_margin = Cm(margin_cm)
    sec.bottom_margin = Cm(margin_cm)
    st = d.styles["Normal"]
    st.font.name, st.font.size = FONT, Pt(base_pt)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    st.paragraph_format.space_after = Pt(4)
    st.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Heading 1", 14, "1F3864"), ("Heading 2", 11.5, "2F5496"), ("Title", 22, "1F3864")):
        h = d.styles[name]
        h.font.name, h.font.size, h.font.bold = FONT, Pt(size), True
        h.font.color.rgb = RGBColor.from_string(color)
        h.element.rPr.rFonts.set(qn("w:ascii"), FONT)
        h.element.rPr.rFonts.set(qn("w:hAnsi"), FONT)
        h.paragraph_format.space_before = Pt(8 if name != "Title" else 0)
        h.paragraph_format.space_after = Pt(3)
        h.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = footer.add_run(f"{title_for_footer}  |  Page ")
    r.font.size = Pt(8.5)
    _field(footer, "PAGE")
    return d


def _field(par, code: str) -> None:
    run = par.add_run()
    run.font.size = Pt(8.5)
    for t, txt in (("begin", None), (None, code), ("end", None)):
        if t:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), t)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


def h1(d, text):
    return d.add_heading(text, level=1)


def h2(d, text):
    return d.add_heading(text, level=2)


def p(d, text: str = "", bold_prefix: str | None = None, size: float | None = None, italic=False, align=None, after=None):
    par = d.add_paragraph()
    if bold_prefix:
        r = par.add_run(bold_prefix)
        r.bold = True
        if size:
            r.font.size = Pt(size)
    r = par.add_run(text)
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if align == "center":
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if after is not None:
        par.paragraph_format.space_after = Pt(after)
    return par


def bullets(d, items, size: float | None = None):
    for it in items:
        par = d.add_paragraph(style="List Bullet")
        par.paragraph_format.space_after = Pt(1.5)
        if isinstance(it, tuple):
            r = par.add_run(it[0])
            r.bold = True
            r2 = par.add_run(it[1])
            if size:
                r.font.size = r2.font.size = Pt(size)
        else:
            r = par.add_run(it)
            if size:
                r.font.size = Pt(size)


def mono(d, text: str, size: float = 7.8):
    par = d.add_paragraph()
    par.paragraph_format.left_indent = Cm(0.3)
    par.paragraph_format.space_after = Pt(3)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "F2F2F2")
    par._p.get_or_add_pPr().append(shd)
    r = par.add_run(text)
    r.font.name, r.font.size = "Consolas", Pt(size)
    return par


def _shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(d, header, rows, widths_cm=None, size: float = 8.5, header_fill="D9E2F3"):
    t = d.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]
        c.text = ""
        r = c.paragraphs[0].add_run(str(h))
        r.bold, r.font.size = True, Pt(size)
        _shade(c, header_fill)
    for row in rows:
        cells = t.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = ""
            r = cells[i].paragraphs[0].add_run(str(v))
            r.font.size = Pt(size)
    for row in t.rows:
        for c in row.cells:
            for par in c.paragraphs:
                par.paragraph_format.space_after = Pt(0)
                par.paragraph_format.line_spacing = 1.0
    if widths_cm:
        t.autofit = False
        for col, w in zip(t.columns, widths_cm):
            col.width = Cm(w)
        for row in t.rows:
            for c, w in zip(row.cells, widths_cm):
                c.width = Cm(w)
    d.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def figure(d, path, width_cm: float, caption: str):
    par = d.add_paragraph()
    par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    par.paragraph_format.keep_with_next = True
    par.add_run().add_picture(str(path), width=Cm(width_cm))
    cap = d.add_paragraph()
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = cap.add_run(caption)
    r.italic, r.font.size = True, Pt(9)


def page_break(d):
    d.add_page_break()
