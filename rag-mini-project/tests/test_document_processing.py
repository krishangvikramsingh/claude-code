import pytest
from docx import Document

from rag.chunking import Chunk, chunk_segments, chunk_text, split_sentences, validate_chunks
from rag.document_loader import Segment, load_document
from rag.errors import EmptyDocumentError, IngestionError, UnsupportedFormatError
from rag.text_cleaner import clean_segments, normalize_text, remove_repeated_lines


# ---- loader ----
def test_load_markdown_sections(tmp_path):
    p = tmp_path / "a.md"
    p.write_text("# Title\n\n## Intro\n\nHello world. This is text.\n\n## Next\n\nMore text here.\n", encoding="utf-8")
    segs = load_document(p)
    assert [s.section for s in segs] == ["Intro", "Next"]
    assert segs[0].page is None


def test_load_docx_headings(tmp_path):
    d = Document()
    d.add_heading("Sec A", 1)
    d.add_paragraph("Body of A.")
    p = tmp_path / "a.docx"
    d.save(p)
    segs = load_document(p)
    assert segs[0].section == "Sec A" and "Body of A." in segs[0].text


def test_load_pdf_pages_from_sample_kb(settings):
    segs = load_document(settings.kb_dir / "01_rag_overview.pdf")
    assert [s.page for s in segs] == [1, 2]


def test_unsupported_format(tmp_path):
    p = tmp_path / "x.xlsx"
    p.write_bytes(b"data")
    with pytest.raises(UnsupportedFormatError):
        load_document(p)


def test_empty_file(tmp_path):
    p = tmp_path / "e.txt"
    p.write_bytes(b"")
    with pytest.raises(EmptyDocumentError):
        load_document(p)


def test_whitespace_only_file(tmp_path):
    p = tmp_path / "w.md"
    p.write_text("   \n\n  \n", encoding="utf-8")
    with pytest.raises(EmptyDocumentError):
        load_document(p)


def test_corrupted_pdf_and_docx(tmp_path):
    for name in ("bad.pdf", "bad.docx"):
        p = tmp_path / name
        p.write_bytes(b"this is not a real file at all")
        with pytest.raises(IngestionError):
            load_document(p)


def test_non_utf8_text(tmp_path):
    p = tmp_path / "b.txt"
    p.write_bytes(b"\xff\xfe\x00bad")
    with pytest.raises(IngestionError):
        load_document(p)


def test_missing_file(tmp_path):
    with pytest.raises(IngestionError):
        load_document(tmp_path / "nope.txt")


# ---- cleaner ----
def test_normalize_whitespace_and_hyphenation():
    t = normalize_text("Hel-\nlo   world\x00 again\nsame para\n\n\n\nNew   para")
    assert t == "Hello world again same para\n\nNew para"


def test_heading_line_kept_separate():
    assert normalize_text("Some Heading\nBody sentence here.") == "Some Heading\n\nBody sentence here."


def test_repeated_header_footer_removed():
    bodies = ["Alpha paragraph content here.", "Beta paragraph content here.", "Gamma paragraph content here."]
    segs = [Segment(f"Header text\n{b}\nPage {i}", page=i) for i, b in zip((1, 2, 3), bodies)]
    out, removed = remove_repeated_lines(segs)
    assert removed >= 6
    assert all("Header text" not in s.text and "Page" not in s.text for s in out)
    assert all("content here" in s.text for s in out)


def test_clean_drops_empty_segments():
    segs, stats = clean_segments([Segment("  \n ", 1), Segment("Real content. Here it is.", 2)])
    assert len(segs) == 1 and stats["empty_segments_dropped"] == 1


# ---- chunking ----
TEXT = " ".join(f"Sentence number {i} talks about topic {i} in some detail." for i in range(40))


def test_chunk_size_respected_and_overlap_present():
    chunks = chunk_text(TEXT, 300, 80)
    assert len(chunks) > 3
    assert all(len(c) <= 300 for c in chunks)
    # overlap: the last sentence of chunk i appears at the start of chunk i+1
    hits = sum(split_sentences(a)[-1] in b for a, b in zip(chunks, chunks[1:]))
    assert hits >= len(chunks) - 2


def test_zero_overlap_has_no_repeats():
    chunks = chunk_text(TEXT, 300, 0)
    joined = " ".join(chunks)
    assert joined.count("Sentence number 7 ") == 1


def test_all_text_covered():
    chunks = chunk_text(TEXT, 300, 80)
    for i in range(40):
        assert any(f"Sentence number {i} " in c for c in chunks)


def test_long_sentence_hard_split():
    chunks = chunk_text("word " * 500, 200, 20)
    assert all(len(c) <= 200 for c in chunks) and len(chunks) > 5


def test_invalid_params():
    with pytest.raises(ValueError):
        chunk_text("abc", 100, 100)
    with pytest.raises(ValueError):
        chunk_text("abc", 0, 0)


def test_metadata_preserved_and_ids_unique():
    segs = [Segment(TEXT, page=3, section="Sec"), Segment(TEXT, page=4, section="Sec2")]
    chunks, _ = chunk_segments("doc.pdf", segs, 300, 50)
    assert {c.page for c in chunks} == {3, 4}
    assert all(c.doc == "doc.pdf" and c.section in ("Sec", "Sec2") for c in chunks)
    assert len({c.chunk_id for c in chunks}) == len(chunks)
    assert validate_chunks(chunks, 300) == []


def test_short_chunks_skipped():
    chunks, st = chunk_segments("d.txt", [Segment("Tiny.", None, None)], 300, 50, min_chars=40)
    assert chunks == [] and st["skipped_short_chunks"] == 1


def test_validate_detects_problems():
    bad = [Chunk("a", "d", None, None, "  ", 2, "h"), Chunk("a", "d", None, None, "x" * 600, 600, "h2")]
    probs = validate_chunks(bad, 500)
    assert any("empty" in p for p in probs) and any("exceeds" in p for p in probs) and any("duplicate" in p for p in probs)
